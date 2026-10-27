"""Service for safely extracting raw tabular data from CSV, XLSX, XLS, and PDF files."""

import csv
from pathlib import Path

import openpyxl
import pandas as pd
import pymupdf
import xlrd

from app.core.config import settings
from app.core.errors import BadRequestError, ValidationError


def detect_csv_encoding(file_path: Path) -> tuple[str, list[str]]:
    """Determine best encoding for CSV and log warnings for non-standard encodings."""
    warnings = []
    with open(file_path, "rb") as f:
        sample = f.read(32768)

    for enc in ("utf-8-sig", "utf-8", "utf-16"):
        try:
            sample.decode(enc)
            return enc, warnings
        except (UnicodeDecodeError, LookupError):
            continue

    # Latin-1 fallback
    warnings.append(
        "CSV file parsed using fallback latin-1 encoding; non-ASCII characters may require review."
    )
    return "latin-1", warnings


def extract_csv(file_path: Path) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Extract tabular data from CSV with encoding detection and delimiter sniffing."""
    encoding, warnings = detect_csv_encoding(file_path)

    # Sniff delimiter
    delimiter = ","
    try:
        with open(file_path, encoding=encoding, errors="replace") as f:
            sample_text = f.read(8192)
            if sample_text:
                sniffer = csv.Sniffer()
                dialect = sniffer.sniff(sample_text, delimiters=",\t;|")
                delimiter = dialect.delimiter
    except Exception:
        delimiter = ","

    try:
        df = pd.read_csv(
            file_path,
            encoding=encoding,
            sep=delimiter,
            nrows=settings.MAX_ROWS + 1,
            dtype=str,
            on_bad_lines="skip",
            low_memory=False,
        )
    except Exception as e:
        raise BadRequestError(f"Failed to parse CSV file: {e!s}")

    if df.empty:
        raise BadRequestError("CSV file contains no data rows.")

    if len(df) > settings.MAX_ROWS:
        raise BadRequestError(f"CSV file exceeds maximum row limit of {settings.MAX_ROWS}.")

    if len(df.columns) > settings.MAX_COLUMNS:
        raise BadRequestError(f"CSV file exceeds maximum column limit of {settings.MAX_COLUMNS}.")

    return df, warnings, []


def extract_xlsx(
    file_path: Path, sheet_name: str | None = None
) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Extract tabular data from XLSX using openpyxl in read-only and data-only mode."""
    warnings: list[str] = []

    try:
        with open(file_path, "rb") as f_in:
            wb = openpyxl.load_workbook(
                filename=f_in,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            available_sheets = list(wb.sheetnames)
            if not available_sheets:
                raise BadRequestError("Excel workbook contains no sheets.")

            if len(available_sheets) > settings.MAX_SHEETS:
                raise BadRequestError(
                    f"Excel workbook contains {len(available_sheets)} sheets, exceeding limit of {settings.MAX_SHEETS}."
                )

            # Choose sheet
            target_sheet_name = sheet_name or available_sheets[0]
            if target_sheet_name not in available_sheets:
                raise BadRequestError(
                    f"Requested sheet '{target_sheet_name}' was not found. Available sheets: {', '.join(available_sheets)}"
                )

            ws = wb[target_sheet_name]
            rows = []
            row_count = 0

            for row in ws.iter_rows(values_only=True):
                # Skip empty lines
                if not any(v is not None and str(v).strip() != "" for v in row):
                    continue
                rows.append(row)
                row_count += 1
                if row_count > (settings.MAX_ROWS + 1):
                    raise BadRequestError(
                        f"Excel sheet '{target_sheet_name}' exceeds maximum row limit of {settings.MAX_ROWS}."
                    )
            wb.close()
    except Exception as e:
        if isinstance(e, BadRequestError):
            raise
        raise BadRequestError(f"Failed to open Excel (.xlsx) workbook: {e!s}")

    if not rows or len(rows) < 2:
        raise BadRequestError(f"Excel sheet '{target_sheet_name}' contains no data rows.")

    header = [str(col).strip() if col is not None else f"col_{i}" for i, col in enumerate(rows[0])]
    if len(header) > settings.MAX_COLUMNS:
        raise BadRequestError(
            f"Excel sheet '{target_sheet_name}' exceeds maximum column limit of {settings.MAX_COLUMNS}."
        )

    data = rows[1:]
    # Normalize row lengths to match header
    normalized_data = [
        [
            str(val).strip() if val is not None else ""
            for val in (list(r) + [""] * (len(header) - len(r)))[: len(header)]
        ]
        for r in data
    ]

    df = pd.DataFrame(normalized_data, columns=header)
    return df, warnings, available_sheets


def extract_xls(
    file_path: Path, sheet_name: str | None = None
) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Extract tabular data from legacy XLS file using xlrd."""
    warnings: list[str] = []

    try:
        book = xlrd.open_workbook(str(file_path), on_demand=True)
    except Exception as e:
        raise BadRequestError(f"Failed to open legacy Excel (.xls) file: {e!s}")

    available_sheets = book.sheet_names()
    if not available_sheets:
        raise BadRequestError("Legacy Excel file contains no sheets.")

    if len(available_sheets) > settings.MAX_SHEETS:
        raise BadRequestError(f"Excel file exceeds maximum sheet limit of {settings.MAX_SHEETS}.")

    target_sheet = sheet_name or available_sheets[0]
    if target_sheet not in available_sheets:
        raise BadRequestError(
            f"Requested sheet '{target_sheet}' not found. Available: {', '.join(available_sheets)}"
        )

    sheet = book.sheet_by_name(target_sheet)
    if sheet.nrows < 2:
        raise BadRequestError(f"Sheet '{target_sheet}' contains insufficient rows.")

    if sheet.nrows > settings.MAX_ROWS + 1:
        raise BadRequestError(f"Sheet '{target_sheet}' exceeds row limit of {settings.MAX_ROWS}.")

    if sheet.ncols > settings.MAX_COLUMNS:
        raise BadRequestError(
            f"Sheet '{target_sheet}' exceeds column limit of {settings.MAX_COLUMNS}."
        )

    header = [str(sheet.cell_value(0, c)).strip() for c in range(sheet.ncols)]
    data = []
    for r in range(1, sheet.nrows):
        row_vals = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
        if any(v != "" for v in row_vals):
            data.append(row_vals)

    df = pd.DataFrame(data, columns=header)
    return df, warnings, available_sheets


def extract_pdf(file_path: Path) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Extract tabular energy data from PDF using PyMuPDF (fitz)."""
    warnings: list[str] = []

    try:
        doc = pymupdf.open(str(file_path))
    except Exception as e:
        raise BadRequestError(f"Failed to open PDF document: {e!s}")

    try:
        # Check text layer presence across all pages
        total_text_length = 0
        all_tables_data: list[list[str]] = []
        found_header: list[str] | None = None

        for page_idx in range(doc.page_count):
            page = doc[page_idx]
            page_text = page.get_text()
            total_text_length += len(page_text.strip())

            # Attempt structured table extraction via page.find_tables()
            try:
                tables = page.find_tables()
                for table in tables:
                    extracted = table.extract()
                    if extracted and len(extracted) >= 2:
                        table_header = [str(c or "").strip() for c in extracted[0]]
                        table_rows = [
                            [str(c or "").strip() for c in row]
                            for row in extracted[1:]
                            if any(str(c or "").strip() for c in row)
                        ]
                        if not found_header:
                            found_header = table_header
                            all_tables_data.extend(table_rows)
                        elif len(table_header) == len(found_header):
                            # Append rows from multi-page continuation table
                            all_tables_data.extend(table_rows)
            except Exception:
                pass

        if total_text_length == 0:
            raise ValidationError(
                "Scanned or image-only PDF detected (no text layer found). "
                "OCR is not supported. Please upload a structured CSV or Excel file instead."
            )

        if not found_header or not all_tables_data:
            raise ValidationError(
                "Could not detect structured tabular energy records with recognizable columns in the PDF. "
                "Please upload a structured CSV or Excel (.xlsx) file instead."
            )

        # Build DataFrame
        df = pd.DataFrame(all_tables_data, columns=found_header)

        if len(df) > settings.MAX_ROWS:
            raise BadRequestError(f"PDF tabular data exceeds row limit of {settings.MAX_ROWS}.")

        warnings.append(
            "Extracted tabular data from PDF. Review column mappings to verify accuracy."
        )
        return df, warnings, []
    finally:
        doc.close()


def extract_data_from_file(
    file_path: Path,
    file_type: str,
    sheet_name: str | None = None,
) -> tuple[pd.DataFrame, list[str], list[str]]:
    """Dispatch file extraction based on validated file type."""
    if file_type == "csv":
        return extract_csv(file_path)
    elif file_type == "xlsx":
        return extract_xlsx(file_path, sheet_name=sheet_name)
    elif file_type == "xls":
        return extract_xls(file_path, sheet_name=sheet_name)
    elif file_type == "pdf":
        return extract_pdf(file_path)
    else:
        raise BadRequestError(f"Unsupported file type: {file_type}")
